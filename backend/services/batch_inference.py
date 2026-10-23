"""Schema validation and safe parsing for CSV/XLSX batch uploads."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from backend.settings import MAX_BATCH_ROWS, MAX_UPLOAD_BYTES, MAX_UPLOAD_COLUMNS
from backend.services.prediction_service import engineer_features, prediction_service


REQUIRED_FEATURES = {
    "AMT_INCOME_TOTAL",
    "AMT_CREDIT",
    "DAYS_BIRTH",
}
METADATA_FIELDS = {"application_id", "name", "SK_ID_CURR"}
TARGET_FIELD = "TARGET"
ALIASES = {
    "application id": "application_id",
    "applicant id": "application_id",
    "applicant name": "name",
}
MAX_VALIDATION_ISSUES = 100
ENGINEERED_FEATURES = {
    "AGE_YEARS",
    "EMPLOYED_YEARS",
    "CREDIT_TO_INCOME",
    "ANNUITY_TO_INCOME",
    "CREDIT_TO_ANNUITY",
    "GOODS_TO_CREDIT",
    "INCOME_PER_FAMILY_MEMBER",
    "INCOME_PER_CHILD",
    "EXT_SOURCE_MEAN",
    "EXT_SOURCE_MAX",
}


class BatchFileError(ValueError):
    """An upload cannot be safely parsed or does not satisfy the file contract."""


@dataclass
class ParsedBatch:
    frame: pd.DataFrame
    file_type: str
    original_columns: list[str]
    normalized_columns: list[str]
    duplicate_columns: list[str]
    unsupported_columns: list[str]
    target_removed: bool
    missing_required_columns: list[str]
    missing_model_columns: list[str]
    row_invalid_columns: dict[int, list[str]]
    row_issues: list[dict[str, Any]]
    row_issue_count: int
    missing_value_counts: dict[str, int]

    @property
    def valid_schema(self) -> bool:
        return not self.duplicate_columns and not self.missing_required_columns


def model_feature_schema() -> tuple[list[str], list[str], set[str]]:
    """Return ordered original feature names and numeric/categorical subsets."""
    ordered_expected = [
        column
        for column in prediction_service.expected_columns
        if column not in ENGINEERED_FEATURES
    ]
    numeric = [
        column
        for name, _, columns in prediction_service.preprocessor.transformers_
        if name == "num"
        for column in columns
        if column not in ENGINEERED_FEATURES
    ]
    categorical = [
        column
        for name, _, columns in prediction_service.preprocessor.transformers_
        if name == "cat"
        for column in columns
        if column not in ENGINEERED_FEATURES
    ]
    if len(ordered_expected) != 120:
        raise RuntimeError(
            f"Saved model input contract expected 120 original features, "
            f"found {len(ordered_expected)}."
        )
    if set(ordered_expected) != set(numeric) | set(categorical):
        raise RuntimeError("Saved preprocessor raw-feature schema is inconsistent.")
    return ordered_expected, numeric, set(categorical)


def canonicalize_column(name: Any, recognized: set[str]) -> str:
    raw = "" if name is None else str(name).strip()
    alias = ALIASES.get(raw.casefold())
    if alias:
        return alias
    by_casefold = {column.casefold(): column for column in recognized}
    return by_casefold.get(raw.casefold(), raw)


def _read_csv(content: bytes) -> tuple[pd.DataFrame, list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise BatchFileError("CSV must be encoded as UTF-8.") from exc
    try:
        header = next(csv.reader(io.StringIO(text)))
    except StopIteration as exc:
        raise BatchFileError("The uploaded CSV is empty.") from exc
    if not header or not any(str(value).strip() for value in header):
        raise BatchFileError("The uploaded CSV has no column headers.")
    if len(header) > MAX_UPLOAD_COLUMNS:
        raise BatchFileError(
            f"Upload exceeds the maximum of {MAX_UPLOAD_COLUMNS:,} columns."
        )
    try:
        frame = pd.read_csv(io.BytesIO(content), nrows=MAX_BATCH_ROWS + 1)
    except (pd.errors.ParserError, pd.errors.EmptyDataError, UnicodeDecodeError) as exc:
        raise BatchFileError("Could not parse the uploaded CSV.") from exc
    return frame, [str(value) for value in header]


def _read_xlsx(content: bytes) -> tuple[pd.DataFrame, list[str]]:
    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=True,
            data_only=True,
            keep_links=False,
        )
    except Exception as exc:
        raise BatchFileError("Could not parse the uploaded XLSX workbook.") from exc
    try:
        worksheet = workbook.active
        if not isinstance(worksheet, ReadOnlyWorksheet):
            raise BatchFileError("The active XLSX sheet must be a worksheet.")
        iterator = worksheet.iter_rows(values_only=True)
        header_values = next(iterator, None)
        if not header_values or not any(
            value is not None and str(value).strip() for value in header_values
        ):
            raise BatchFileError("The uploaded worksheet has no column headers.")
        if worksheet.max_column > MAX_UPLOAD_COLUMNS:
            raise BatchFileError(
                f"Upload exceeds the maximum of {MAX_UPLOAD_COLUMNS:,} columns."
            )
        headers = ["" if value is None else str(value) for value in header_values]
        rows = []
        for row_number, values in enumerate(iterator, start=2):
            if row_number > MAX_BATCH_ROWS + 1:
                raise BatchFileError(
                    f"Upload exceeds the maximum of {MAX_BATCH_ROWS:,} rows."
                )
            rows.append(values)
        frame = pd.DataFrame(rows, columns=headers)
    except BatchFileError:
        raise
    except Exception as exc:
        raise BatchFileError("Could not read the uploaded worksheet.") from exc
    finally:
        workbook.close()
    return frame, headers


def parse_batch_upload(filename: str | None, content: bytes) -> ParsedBatch:
    if not filename:
        raise BatchFileError("A filename is required.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise BatchFileError(
            f"Upload exceeds the maximum size of {MAX_UPLOAD_BYTES // (1024 * 1024)} MB."
        )
    extension = filename.rsplit(".", maxsplit=1)[-1].casefold()
    if extension == "csv":
        frame, original_columns = _read_csv(content)
        file_type = "csv"
    elif extension == "xlsx":
        frame, original_columns = _read_xlsx(content)
        file_type = "xlsx"
    else:
        raise BatchFileError("Only .csv and .xlsx files are supported.")
    if len(frame) > MAX_BATCH_ROWS:
        raise BatchFileError(f"Upload exceeds the maximum of {MAX_BATCH_ROWS:,} rows.")
    if frame.empty:
        raise BatchFileError("The uploaded file contains no applicant rows.")

    model_columns, numeric_columns, categorical_columns = model_feature_schema()
    recognized = set(model_columns) | METADATA_FIELDS | {TARGET_FIELD}
    normalized = [
        canonicalize_column(column, recognized) for column in original_columns
    ]
    canonical_counts = pd.Series(normalized).value_counts()
    duplicates = sorted(canonical_counts[canonical_counts > 1].index.tolist())
    # Prevent Pandas' duplicate-header suffixes from being misinterpreted as
    # legitimate arbitrary columns.
    duplicates.extend(
        sorted(
            column
            for column in normalized
            if re.search(r"\.\d+$", column) and column.split(".")[0] in recognized
        )
    )
    duplicate_columns = sorted(set(duplicates))

    # Assign normalized names positionally only after duplicate detection.
    frame.columns = normalized
    target_removed = TARGET_FIELD in frame.columns
    unsupported_columns = sorted(set(frame.columns) - recognized)
    supported_columns = [
        column
        for column in frame.columns
        if column in model_columns or column in METADATA_FIELDS
    ]
    model_frame = frame.loc[:, supported_columns].copy()
    missing_required = sorted(REQUIRED_FEATURES - set(model_frame.columns))
    missing_model_columns = sorted(set(model_columns) - set(model_frame.columns))

    for column in numeric_columns:
        if column not in model_frame:
            continue
        values = model_frame[column]
        converted = pd.to_numeric(values, errors="coerce")
        invalid = values.notna() & (
            converted.isna() | ~np.isfinite(converted)
        )
        if column == "DAYS_BIRTH":
            invalid |= converted.notna() & (converted >= 0)
        if column in {"AMT_INCOME_TOTAL", "AMT_CREDIT"}:
            invalid |= converted.notna() & (converted <= 0)
        model_frame[column] = converted

    row_invalid_columns: dict[int, list[str]] = {}
    for row_position in range(len(model_frame)):
        invalid_columns: list[str] = []
        for column in numeric_columns:
            if column not in model_frame:
                continue
            value = model_frame.iloc[row_position][column]
            original_value = frame.iloc[row_position][column]
            required_missing = (
                column in REQUIRED_FEATURES and pd.isna(original_value)
            )
            invalid_type = (
                pd.notna(original_value)
                and pd.isna(value)
            )
            invalid_domain = (
                pd.notna(value)
                and not np.isfinite(value)
            ) or (
                column == "DAYS_BIRTH" and pd.notna(value) and value >= 0
            ) or (
                column in {"AMT_INCOME_TOTAL", "AMT_CREDIT"}
                and pd.notna(value) and value <= 0
            )
            if required_missing or invalid_type or invalid_domain:
                invalid_columns.append(column)
        for column in categorical_columns:
            if column in model_frame:
                value = model_frame.iloc[row_position][column]
                if pd.notna(value) and not isinstance(value, str):
                    invalid_columns.append(column)
        if invalid_columns:
            row_invalid_columns[row_position] = invalid_columns
    row_issues = [
        {
            "row_number": row_position + 2,
            "columns": columns,
            "code": "invalid_value",
            "message": "One or more required/numeric fields are missing or invalid.",
        }
        for row_position, columns in list(row_invalid_columns.items())[
            :MAX_VALIDATION_ISSUES
        ]
    ]
    row_issue_count = len(row_invalid_columns)
    missing_value_counts = {
        column: int(model_frame[column].isna().sum())
        for column in model_columns
        if column in model_frame and model_frame[column].isna().any()
    }

    return ParsedBatch(
        frame=model_frame,
        file_type=file_type,
        original_columns=original_columns,
        normalized_columns=normalized,
        duplicate_columns=duplicate_columns,
        unsupported_columns=unsupported_columns,
        target_removed=target_removed,
        missing_required_columns=missing_required,
        missing_model_columns=missing_model_columns,
        row_invalid_columns=row_invalid_columns,
        row_issues=row_issues,
        row_issue_count=row_issue_count,
        missing_value_counts=missing_value_counts,
    )


def validation_payload(parsed: ParsedBatch) -> dict[str, Any]:
    model_columns, _, _ = model_feature_schema()
    included = [column for column in model_columns if column in parsed.frame]
    return {
        "valid": parsed.valid_schema and parsed.row_issue_count == 0,
        "file_type": parsed.file_type,
        "row_count": len(parsed.frame),
        "recognized_model_columns": included,
        "recognized_model_column_count": len(included),
        "expected_model_column_count": len(model_columns),
        "missing_required_columns": parsed.missing_required_columns,
        "missing_optional_model_columns": parsed.missing_model_columns,
        "duplicate_columns": parsed.duplicate_columns,
        "unsupported_columns": parsed.unsupported_columns,
        "target_removed": parsed.target_removed,
        "row_issue_count": parsed.row_issue_count,
        "row_issues": parsed.row_issues,
        "row_issues_truncated": parsed.row_issue_count > len(parsed.row_issues),
        "missing_value_counts": parsed.missing_value_counts,
        "warnings": [
            warning
            for condition, warning in (
                (
                    parsed.missing_model_columns,
                    "Some model columns are absent; the saved training-time "
                    "imputer will be used. This is partial-feature inference.",
                ),
                (
                    parsed.unsupported_columns,
                    "Unsupported columns will be quarantined and excluded from inference.",
                ),
                (
                    parsed.target_removed,
                    "TARGET was excluded from model inputs.",
                ),
            )
            if condition
        ],
    }


def score_batch(parsed: ParsedBatch, threshold: float) -> dict[str, Any]:
    if not parsed.valid_schema:
        raise BatchFileError(
            "Required batch schema is invalid: "
            + ", ".join(parsed.missing_required_columns + parsed.duplicate_columns)
        )
    model_columns, _, _ = model_feature_schema()
    results: list[dict[str, Any] | None] = [None] * len(parsed.frame)
    valid_positions: list[int] = []
    valid_rows: list[dict[str, Any]] = []
    for position, (_, row) in enumerate(parsed.frame.iterrows()):
        row_number = position + 2
        application_id_value = row.get("application_id")
        if pd.isna(application_id_value) or not str(application_id_value).strip():
            application_id_value = row.get("SK_ID_CURR")
        if pd.isna(application_id_value) or not str(application_id_value).strip():
            application_id_value = f"ROW-{position + 1:06d}"
        application_id = str(application_id_value)

        result: dict[str, Any] = {
            "row_number": row_number,
            "application_id": application_id,
            "risk_probability": None,
            "predicted_class": None,
            "provisional_classification": None,
            "classification_threshold": threshold,
            "feature_coverage": 0.0,
            "missing_model_features": len(model_columns),
            "status": "ERROR",
            "error": None,
        }
        invalid_columns = parsed.row_invalid_columns.get(position, [])
        if invalid_columns:
            result["error"] = (
                "Invalid value(s) in: " + ", ".join(sorted(set(invalid_columns)))
            )
            results[position] = result
            continue

        applicant = {
            column: (None if pd.isna(value) else value)
            for column, value in row.items()
            if column in model_columns
        }
        provided_count = sum(
            column in applicant and applicant[column] is not None
            for column in model_columns
        )
        result["feature_coverage"] = round(provided_count / len(model_columns), 4)
        result["missing_model_features"] = len(model_columns) - provided_count
        valid_positions.append(position)
        valid_rows.append(
            {
                column: applicant.get(column, np.nan)
                for column in prediction_service.expected_columns
            }
        )
        results[position] = result

    if valid_rows:
        feature_frame = pd.DataFrame(
            valid_rows,
            columns=prediction_service.expected_columns,
        )
        if "DAYS_EMPLOYED" in feature_frame.columns:
            feature_frame["DAYS_EMPLOYED"] = feature_frame[
                "DAYS_EMPLOYED"
            ].replace(365243, np.nan)
        engineered_frame = engineer_features(feature_frame)
        processed = prediction_service.preprocessor.transform(
            engineered_frame.loc[:, prediction_service.expected_columns]
        )
        probabilities = prediction_service.model.predict_proba(processed)[:, 1]
        if len(probabilities) != len(valid_positions):
            raise RuntimeError("Batch prediction count did not match valid input rows.")

        for position, probability in zip(valid_positions, probabilities):
            result_entry = results[position]
            if result_entry is None:
                raise RuntimeError("A valid batch row has no aligned result record.")
            predicted_class = int(probability >= threshold)
            result_entry.update(
                {
                    "risk_probability": round(float(probability), 6),
                    "predicted_class": predicted_class,
                    "provisional_classification": (
                        "elevated_default_risk"
                        if predicted_class == 1
                        else "lower_default_risk"
                    ),
                    "status": "scored",
                }
            )

    aligned_results = [
        result
        for result in results
        if result is not None
    ]
    if len(aligned_results) != len(parsed.frame):
        raise RuntimeError("Batch result count did not match uploaded rows.")
    successful = sum(result["status"] == "scored" for result in aligned_results)
    return {
        "batch_id": None,
        "total_applications": len(aligned_results),
        "successful_predictions": successful,
        "failed_predictions": len(aligned_results) - successful,
        "classification_threshold": threshold,
        "feature_warnings": validation_payload(parsed)["warnings"],
        "schema_validation": validation_payload(parsed),
        "results": aligned_results,
    }
